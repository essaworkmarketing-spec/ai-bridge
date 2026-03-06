from fastapi import FastAPI, UploadFile, Form, HTTPException
from fastapi.middleware.cors import CORSMiddleware
import httpx
import base64
import os
import json
import re
import logging
import io

# ── Logging ──────────────────────────────────────────────────────────────────
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

# ── App ───────────────────────────────────────────────────────────────────────
app = FastAPI(title="Umrah Rate Sheet Extractor API", version="1.0.0")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)

# ── Config ────────────────────────────────────────────────────────────────────
OPENROUTER_API_KEY = os.getenv("OPENROUTER_API_KEY", "")
OPENROUTER_URL = "https://openrouter.ai/api/v1/chat/completions"
MODEL = os.getenv("MODEL", "openai/gpt-4o")

# ── Prompt ────────────────────────────────────────────────────────────────────
SYSTEM_PROMPT = """
You are a strict structured data extraction engine for Umrah/Hajj rate sheets.

You will receive data in ANY format: image, PDF, plain text, or Excel (converted to text). Extract from all formats equally.

Return ONLY a single raw JSON object. No explanation. No markdown. No comments. No text outside JSON.

The JSON must ALWAYS contain exactly these six top-level keys:
{
  "hotels": [],
  "vehicles": [],
  "pvt_transport": [],
  "sharing_transport": [],
  "pvt_ziyarah": [],
  "sharing_ziyarah": []
}

════════════════════════════════════════
CUSTOM INSTRUCTIONS:
════════════════════════════════════════
If the user provides additional instructions, follow them strictly.
Custom instructions override default behavior where they conflict.
Apply them during extraction, not after.

════════════════════════════════════════
HOTELS EXTRACTION RULES:
════════════════════════════════════════
Each hotel object:
{
  "city": "Makkah" or "Madinah",
  "name": "full hotel name as written",
  "location": "location/area as written",
  "distance": "distance value as written e.g. 700-800 MTR or SHUTTLE SERVICE",
  "sharing": null or number,
  "quint": null or number,
  "quad": null or number,
  "triple": null or number,
  "double": null or number,
  "flat_room_rate": null or number
}

- Extract EVERY row from BOTH Makkah and Madinah hotel tables.
- CRITICAL: If hotel is under "MAKKAH HOTELS" header set city = "Makkah". If under "MADINAH HOTELS" header set city = "Madinah". Never leave city blank or unknown.
- If a rate cell says "N/A" or is blank, use null.
- If a cell says "FLAT ROOM RATE 700/-" extract 700 into flat_room_rate, set other rates to null.
- Preserve exact numbers. Never skip any row or rate column.

════════════════════════════════════════
PRIVATE TRANSPORT EXTRACTION RULES:
════════════════════════════════════════
pvt_transport = per vehicle pricing. Input is usually a table with vehicle columns.

Each object:
{
  "route": "route name as written",
  "camry_sonata": null or number,
  "h1_hyundai": null or number,
  "gmc": null or number,
  "hiace": null or number,
  "coaster": null or number,
  "grand_cabin": null or number,
  "bus": null or number
}

- Include EVERY route row.
- Match column headers to vehicle fields exactly.
- If a vehicle column does not exist in the data, use null.

════════════════════════════════════════
SHARING TRANSPORT EXTRACTION RULES:
════════════════════════════════════════
sharing_transport = per person transport pricing between cities/routes.

Input can be ANY of these formats:
  FORMAT A (table): route rows with per_person price column
  FORMAT B (plain text bullet list):
    - Jeddah to Makkah 70sar
    - Makkah to Madinah 70 sar
  FORMAT C (labeled plain text):
    sharing transport rates
    - Jeddah to Makkah 70sar
    - Makkah to Jeddah 70 sar

Each object:
{
  "route": "route name e.g. Jeddah to Makkah",
  "per_person": number,
  "currency": "SAR"
}

- Parse ALL formats above equally.
- Strip sar/SAR/SR suffix, extract only the number into per_person.
- Extract EVERY route without skipping.
- If a header says "sharing transport" treat ALL items below it as sharing_transport.

════════════════════════════════════════
SHARING ZIYARAH EXTRACTION RULES:
════════════════════════════════════════
sharing_ziyarah = per person ziyarah tour pricing.

Input can be ANY of these formats:
  FORMAT A (table): tour rows with per_person price column
  FORMAT B (single line per tour):
    Badar Ziyarat SAR 70/person
    Madinah Ziyarat SAR 70/person
  FORMAT C (alternating lines):
    Badar Ziyarat
    SAR 70/person
    Madinah Ziyarat
    SAR 70/person

Each object:
{
  "tour_name": "name as written",
  "city": "Makkah" or "Madinah",
  "per_person": number,
  "currency": "SAR"
}

City detection rules:
- Badar Ziyarat = "Madinah"
- Madinah Ziyarat = "Madinah"
- Makkah Ziyarat = "Makkah"
- Taif Ziyarat = "Makkah"
- Any tour with "Madinah" in name = "Madinah"
- Any tour with "Makkah" or "Taif" in name = "Makkah"

- Parse ALL formats above equally.
- Strip SAR/ prefix and /person suffix, extract only the number.
- Extract EVERY tour without skipping.

════════════════════════════════════════
PRIVATE ZIYARAH EXTRACTION RULES:
════════════════════════════════════════
pvt_ziyarah = per vehicle ziyarah pricing.

Each object:
{
  "city": "Makkah" or "Madinah",
  "tour_name": "name as written",
  "camry_sonata": null or number,
  "h1_hyundai": null or number,
  "gmc": null or number,
  "hiace": null or number,
  "coaster": null or number,
  "per_person": null or number
}

════════════════════════════════════════
VEHICLES EXTRACTION RULES:
════════════════════════════════════════
If there is a standalone vehicle list or capacity table:
{
  "type": "vehicle type",
  "capacity": null or number,
  "description": "any notes"
}

════════════════════════════════════════
GLOBAL RULES:
════════════════════════════════════════
- Never guess or invent values.
- Preserve all numbers exactly as written.
- Extract EVERY row from EVERY table and EVERY item from EVERY list — zero skipping.
- If a section is not present in the input, return empty array [].
- Handle image, PDF, plain text, and Excel (as text) inputs equally well.
- Do NOT wrap output in markdown fences or any extra text.
- Output must be valid parseable JSON only.
"""

# ── File type helpers ─────────────────────────────────────────────────────────
IMAGE_EXTENSIONS = {"jpg", "jpeg", "png", "webp", "gif"}
TEXT_EXTENSIONS = {"txt", "csv", "text"}
PDF_EXTENSIONS = {"pdf"}
EXCEL_EXTENSIONS = {"xlsx", "xls"}

MIME_MAP = {
    "jpg": "image/jpeg",
    "jpeg": "image/jpeg",
    "png": "image/png",
    "webp": "image/webp",
    "gif": "image/gif",
    "pdf": "application/pdf",
}

REQUIRED_KEYS = [
    "hotels",
    "vehicles",
    "pvt_transport",
    "sharing_transport",
    "pvt_ziyarah",
    "sharing_ziyarah",
]


def get_file_ext(filename: str) -> str:
    return (filename or "").lower().rsplit(".", 1)[-1]


def detect_mime_type(filename: str) -> str:
    ext = get_file_ext(filename)
    return MIME_MAP.get(ext, "image/jpeg")


def get_file_category(filename: str) -> str:
    ext = get_file_ext(filename)
    if ext in IMAGE_EXTENSIONS:
        return "image"
    if ext in PDF_EXTENSIONS:
        return "pdf"
    if ext in TEXT_EXTENSIONS:
        return "text"
    if ext in EXCEL_EXTENSIONS:
        return "excel"
    return "image"


def excel_to_text(file_bytes: bytes) -> str:
    """Convert Excel file to plain text for extraction."""
    try:
        import openpyxl
        workbook = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
        lines = []
        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            lines.append(f"=== Sheet: {sheet_name} ===")
            for row in sheet.iter_rows(values_only=True):
                # Filter out completely empty rows
                row_values = [str(cell) if cell is not None else "" for cell in row]
                if any(v.strip() for v in row_values):
                    lines.append("\t".join(row_values))
        return "\n".join(lines)
    except Exception as e:
        raise ValueError(f"Could not read Excel file: {str(e)}")


def safe_json_parse(text: str) -> dict:
    text = text.strip()
    text = re.sub(r"^```(?:json)?\s*", "", text, flags=re.IGNORECASE)
    text = re.sub(r"\s*```$", "", text.strip())
    text = text.strip()

    try:
        return json.loads(text)
    except json.JSONDecodeError:
        pass

    start = text.find("{")
    end = text.rfind("}") + 1
    if start == -1 or end == 0:
        raise ValueError("No JSON object found in model response")

    cleaned = text[start:end]
    cleaned = re.sub(r",\s*}", "}", cleaned)
    cleaned = re.sub(r",\s*]", "]", cleaned)

    return json.loads(cleaned)


def ensure_schema(data: dict) -> dict:
    for key in REQUIRED_KEYS:
        if key not in data or not isinstance(data[key], list):
            data[key] = []
    return data


def build_message_content(
    file_bytes: bytes,
    filename: str,
    instructions: str,
) -> list:
    category = get_file_category(filename)

    instruction_block = ""
    if instructions.strip():
        instruction_block = f"\n\nCUSTOM INSTRUCTIONS (follow these strictly during extraction):\n{instructions.strip()}"

    extraction_prompt = f"Extract all structured pricing data from this rate sheet. Follow all extraction rules strictly.{instruction_block}"

    # ── Plain text ────────────────────────────────────────────────────────────
    if category == "text":
        try:
            text_content = file_bytes.decode("utf-8")
        except UnicodeDecodeError:
            text_content = file_bytes.decode("latin-1")

        return [
            {
                "type": "text",
                "text": f"{extraction_prompt}\n\nRATE SHEET CONTENT:\n{text_content}",
            }
        ]

    # ── Excel ─────────────────────────────────────────────────────────────────
    if category == "excel":
        text_content = excel_to_text(file_bytes)
        return [
            {
                "type": "text",
                "text": f"{extraction_prompt}\n\nRATE SHEET CONTENT (from Excel):\n{text_content}",
            }
        ]

    # ── Image or PDF ──────────────────────────────────────────────────────────
    image_base64 = base64.b64encode(file_bytes).decode("utf-8")
    mime_type = detect_mime_type(filename)

    return [
        {
            "type": "image_url",
            "image_url": {
                "url": f"data:{mime_type};base64,{image_base64}"
            },
        },
        {
            "type": "text",
            "text": extraction_prompt,
        },
    ]


# ── Routes ────────────────────────────────────────────────────────────────────

@app.get("/")
async def root():
    return {
        "status": "ok",
        "service": "Umrah Rate Sheet Extractor",
        "model": MODEL,
        "version": "1.0.0",
        "supported_formats": ["jpg", "jpeg", "png", "webp", "gif", "pdf", "txt", "csv", "xlsx", "xls"],
    }


@app.get("/health")
async def health():
    return {"status": "healthy"}


@app.post("/extract")
async def extract_data(
    file: UploadFile,
    instructions: str = Form(default=""),
):
    if not OPENROUTER_API_KEY:
        raise HTTPException(
            status_code=500,
            detail="OPENROUTER_API_KEY environment variable is not set.",
        )

    file_bytes = await file.read()
    if not file_bytes:
        raise HTTPException(status_code=400, detail="Uploaded file is empty.")

    filename = file.filename or "upload.jpg"
    category = get_file_category(filename)

    logger.info(f"Processing: {filename} | type: {category} | size: {len(file_bytes)} bytes")
    if instructions.strip():
        logger.info(f"Custom instructions provided: {instructions.strip()[:200]}")

    try:
        user_content = build_message_content(file_bytes, filename, instructions)
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))

    payload = {
        "model": MODEL,
        "temperature": 0.1,
        "max_tokens": 4000,
        "messages": [
            {"role": "system", "content": SYSTEM_PROMPT},
            {"role": "user", "content": user_content},
        ],
    }

    headers = {
        "Authorization": f"Bearer {OPENROUTER_API_KEY}",
        "Content-Type": "application/json",
        "HTTP-Referer": "https://umrah-extractor.onrender.com",
        "X-Title": "Umrah Rate Sheet Extractor",
    }

    try:
        async with httpx.AsyncClient(timeout=120.0) as client:
            response = await client.post(OPENROUTER_URL, headers=headers, json=payload)
    except httpx.TimeoutException:
        raise HTTPException(status_code=504, detail="OpenRouter request timed out. Try again.")
    except httpx.RequestError as e:
        raise HTTPException(status_code=502, detail=f"Network error: {str(e)}")

    if response.status_code != 200:
        logger.error(f"OpenRouter error {response.status_code}: {response.text[:500]}")
        raise HTTPException(
            status_code=502,
            detail=f"OpenRouter returned {response.status_code}: {response.text[:300]}",
        )

    try:
        result = response.json()
        content = result["choices"][0]["message"]["content"]
    except (KeyError, IndexError, json.JSONDecodeError) as e:
        raise HTTPException(status_code=502, detail=f"Unexpected response from OpenRouter: {str(e)}")

    logger.info(f"Model response (first 300 chars): {content[:300]}")

    try:
        parsed = safe_json_parse(content)
    except Exception as e:
        logger.error(f"JSON parse failed: {content[:1000]}")
        raise HTTPException(
            status_code=422,
            detail={
                "error": "Model returned invalid JSON. Try again.",
                "details": str(e),
                "raw_response": content[:1000],
            },
        )

    parsed = ensure_schema(parsed)
    logger.info("Extraction successful.")
    return parsed
